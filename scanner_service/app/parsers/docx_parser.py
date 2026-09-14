import io
import zipfile
import re
from typing import Dict, Any, List
from PIL import Image

try:
    import docx
except ImportError:
    docx = None

try:
    import pytesseract
except ImportError:
    pytesseract = None

class DOCXParser:
    """
    Parses Word Documents (.docx):
    - Extracts headings, paragraphs, and tables using python-docx.
    - Unpacks word/media/ archive to extract embedded images (e.g. certificates, badges)
      and runs OCR on them, merging results into the output payload.
    """
    
    @staticmethod
    def parse(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        if docx is None:
            raise RuntimeError("python-docx is required for DOCX parsing.")
        
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = []
        tables = []
        fields = []
        embedded_ocr_text = []
        
        # 1. Extract Paragraphs and Headings
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append({
                    "style": p.style.name if p.style else "Normal",
                    "text": text
                })
        
        # 2. Extract Document Tables
        for t_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(row_cells)
            
            if not table_rows:
                continue
            
            headers = table_rows[0]
            data_rows = table_rows[1:] if len(table_rows) > 1 else []
            
            tables.append({
                "table_index": t_idx + 1,
                "headers": headers,
                "row_count": len(data_rows),
                "rows": data_rows
            })
            
            # Form field entries from table headers & columns
            for col_idx, header in enumerate(headers):
                if not header:
                    header = f"Column_{col_idx+1}"
                
                col_values = []
                is_numeric = True
                for row in data_rows:
                    if col_idx < len(row):
                        v = row[col_idx]
                        col_values.append(v)
                        # Check numeric
                        clean_v = v.replace(',', '').replace('%', '').strip()
                        if clean_v:
                            try:
                                float(clean_v)
                            except ValueError:
                                is_numeric = False
                        else:
                            is_numeric = False
                
                fields.append({
                    "table_index": t_idx + 1,
                    "name": header,
                    "type": "numeric" if is_numeric and col_values else "categorical",
                    "unit": "%" if "%" in header else None,
                    "sample_values": col_values[:5],
                    "distinct_count": len(set(col_values)),
                    "total_count": len(col_values)
                })

        # 3. Unpack embedded images from word/media/ ZIP archive
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes), 'r') as z:
                media_files = [f for f in z.namelist() if f.startswith('word/media/')]
                for mf in media_files:
                    img_data = z.read(mf)
                    if pytesseract is not None and len(img_data) > 1024:
                        try:
                            img = Image.open(io.BytesIO(img_data))
                            ocr_text = pytesseract.image_to_string(img).strip()
                            if len(ocr_text) > 8:
                                embedded_ocr_text.append({
                                    "media_path": mf,
                                    "type": "docx_embedded_image_ocr",
                                    "text": ocr_text
                                })
                        except Exception:
                            pass
        except Exception:
            pass

        # 4. Extract KPI stat patterns from paragraph text
        combined_text = "\n".join([p["text"] for p in paragraphs])
        stat_patterns = re.findall(r'([A-Za-z\s\(\)\-\/]{3,40})\s*[:\-\=]\s*([0-9\.,]+)\s*([%\w]*)', combined_text)
        for name, val_str, unit in stat_patterns:
            try:
                clean_val = float(val_str.replace(',', ''))
                fields.append({
                    "name": name.strip(),
                    "type": "numeric",
                    "unit": unit.strip() if unit else None,
                    "sample_values": [clean_val],
                    "distinct_count": 1,
                    "total_count": 1
                })
            except ValueError:
                pass

        return {
            "format": "docx",
            "filename": filename,
            "paragraphs_count": len(paragraphs),
            "tables": tables,
            "fields": fields,
            "raw_text": combined_text,
            "embedded_ocr_text": embedded_ocr_text
        }
