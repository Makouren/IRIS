"""
Tests for DocxExtractor.
"""

from pathlib import Path
from docx import Document
import pytest

from app.extractors.docx import DocxExtractor
from app.structurer import structure_extraction


def test_docx_extraction(tmp_path: Path):
    """Test extracting paragraphs and tables from a Word document."""
    docx_path = tmp_path / "test_report.docx"

    doc = Document()
    doc.add_heading("Institutional QA Evaluation", level=1)
    doc.add_paragraph("Passing Rate: 91.5%")
    doc.add_paragraph("Employment Rate: 84.0%")

    # Add a table
    table = doc.add_table(rows=3, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Department"
    hdr_cells[1].text = "Score"

    row1 = table.rows[1].cells
    row1[0].text = "Computer Science"
    row1[1].text = "94"

    row2 = table.rows[2].cells
    row2[0].text = "Information Technology"
    row2[1].text = "89"

    doc.save(str(docx_path))

    extractor = DocxExtractor()
    result = extractor.extract(docx_path)

    assert result.metadata["file_type"] == "docx"
    assert result.metadata["paragraph_count"] >= 2
    assert result.metadata["table_count"] == 1
    assert len(result.tables) == 1

    fields = structure_extraction(result)
    assert len(fields) >= 2
